from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import re

ROOT = Path(__file__).parent
LAB = ROOT / 'material_practico'

def new_doc(title, subtitle):
    d = Document()
    s = d.sections[0]
    s.page_width, s.page_height = Cm(21), Cm(29.7)
    s.top_margin = s.bottom_margin = Cm(2)
    s.left_margin = s.right_margin = Cm(2.1)
    s.footer_distance = Cm(0.9)
    styles = d.styles
    for border in styles.element.xpath('.//w:pBdr'):
        border.getparent().remove(border)
    for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Heading 3']:
        styles[name].font.name = 'Calibri'
        styles[name].font.color.rgb = RGBColor(0,0,0)
    normal = styles['Normal']
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.08
    for name, size in [('Title', 25), ('Subtitle', 13), ('Heading 1', 17), ('Heading 2', 13), ('Heading 3', 11)]:
        styles[name].font.size = Pt(size)
        styles[name].paragraph_format.space_before = Pt(12)
        styles[name].paragraph_format.space_after = Pt(7)
    styles['Title'].font.bold = True
    styles['Subtitle'].font.italic = False
    code = styles.add_style('Codigo', 1)
    code.font.name, code.font.size = 'Liberation Mono', Pt(8.5)
    code.paragraph_format.line_spacing = 1.0
    code.paragraph_format.space_after = Pt(0)
    code.paragraph_format.space_before = Pt(0)
    code.paragraph_format.widow_control = False
    d.core_properties.title = title
    d.core_properties.subject = subtitle
    d.core_properties.author = ''
    d.add_paragraph(title, 'Title')
    d.add_paragraph(subtitle, 'Subtitle')
    footer = s.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = footer.add_run('Big Data  |  ')
    r.font.size = Pt(9)
    field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE')
    footer._p.append(field)
    return d

def p(d, text):
    return d.add_paragraph(text)

def h(d, text, level=1):
    d.add_heading(text, level)

def text(d, value):
    for block in value.strip().split('\n\n'):
        block = block.strip()
        if not block: continue
        if block.startswith('### '): h(d, block[4:], 3)
        elif block.startswith('## '): h(d, block[3:], 2)
        elif block.startswith('# '): h(d, block[2:], 1)
        elif block.startswith('- '):
            for line in block.splitlines():
                d.add_paragraph(line[2:], 'List Bullet')
        else: p(d, ' '.join(block.splitlines()))

def table(d, headers, rows, widths=None):
    t=d.add_table(rows=1, cols=len(headers))
    t.alignment=WD_TABLE_ALIGNMENT.CENTER
    t.autofit=False
    if widths:
        for col,w in zip(t.columns,widths): col.width=Cm(w)
    for c,v in zip(t.rows[0].cells,headers): c.text=v
    for row in rows:
        for c,v in zip(t.add_row().cells,row): c.text=str(v)
    for i,row in enumerate(t.rows):
        pr=row._tr.get_or_add_trPr()
        cant=OxmlElement('w:cantSplit');pr.append(cant)
        if i==0:
            repeat=OxmlElement('w:tblHeader');pr.append(repeat)
        for j,c in enumerate(row.cells):
            if widths: c.width=Cm(widths[j])
            c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr=c._tc.get_or_add_tcPr()
            shade=OxmlElement('w:shd');shade.set(qn('w:fill'), 'DCE6EF' if i==0 else 'FFFFFF');tcpr.append(shade)
            borders=OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right']:
                e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
            tcpr.append(borders)
            margins=OxmlElement('w:tcMar')
            for edge in ['top','left','bottom','right']:
                e=OxmlElement('w:'+edge);e.set(qn('w:w'),'90');e.set(qn('w:type'),'dxa');margins.append(e)
            tcpr.append(margins)
            for para in c.paragraphs:
                if headers[j] in ['Código','Sincrónico','Autónomo','Total',
                                  'Peso','Peso interno','Lote','Pedidos válidos',
                                  'Importe en centavos','Valor esperado']:
                    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                para.paragraph_format.space_after=Pt(3)
                para.paragraph_format.space_before=Pt(3)
                para.paragraph_format.line_spacing=1.0
                for run in para.runs:
                    run.font.size=Pt(10)
                    run.font.bold=(i==0)
    d.add_paragraph().paragraph_format.space_after=Pt(1)
    return t

def code(d, s):
    for line in s.strip('\n').splitlines():
        d.add_paragraph(line, 'Codigo')
    d.add_paragraph().paragraph_format.space_after=Pt(2)

def labcode(d, name):
    code(d, (LAB/name).read_text())

def chapter(d, title):
    h(d,title)

REFERENCES = [
('R1', 'Apache Software Foundation', '2025', 'PySpark 4.0.1 Installation', 'https://spark.apache.org/docs/4.0.1/api/python/getting_started/install.html'),
('R2', 'Apache Software Foundation', '2025', 'RDD Programming Guide Spark 4.0.1', 'https://spark.apache.org/docs/4.0.1/rdd-programming-guide.html'),
('R3', 'Apache Software Foundation', '2025', 'Performance Tuning Spark 4.0.1', 'https://spark.apache.org/docs/4.0.1/sql-performance-tuning.html'),
('R4', 'Apache Software Foundation', '2025', 'Structured Streaming Programming Guide Spark 4.0.1', 'https://spark.apache.org/docs/4.0.1/streaming/apis-on-dataframes-and-datasets.html'),
('R5', 'DuckDB Foundation', 's. f.', 'Python API', 'https://duckdb.org/docs/current/clients/python/overview'),
('R6', 'DuckDB Foundation', 's. f.', 'Tuning Workloads', 'https://duckdb.org/docs/current/guides/performance/how_to_tune_workloads'),
('R7', 'Apache Software Foundation', 's. f.', 'Apache Parquet Overview', 'https://parquet.apache.org/docs/overview/'),
('R8', 'scikit-learn developers', 's. f.', 'Common pitfalls and recommended practices', 'https://scikit-learn.org/stable/common_pitfalls.html'),
]

def references(d):
    p(d,'Fuentes técnicas de consulta. Los códigos R1 a R8 remiten a estas referencias. Consulta realizada el 29 de septiembre de 2026. Las actividades, los casos y los datos sintéticos del curso son elaboraciones didácticas; no representan observaciones de una organización real.')
    for key,author,year,title,url in REFERENCES:
        p(d, f'[{key}] {author}. ({year}). {title}. {url}')
    p(d,'Lecturas de profundización sugeridas: Kleppmann, M. (2017). Designing Data-Intensive Applications. O’Reilly Media. Leskovec, J., Rajaraman, A. y Ullman, J. D. (2020). Mining of Massive Datasets, tercera edición. Cambridge University Press. Estas obras amplían el curso y no son necesarias para ejecutar las prácticas.')

plan = new_doc('Programa académico de Big Data', 'Fundamentos ingeniería y analítica aplicada\nPosgrado  |  Modalidad online  |  64 horas')
text(plan, '''Este programa desarrolla competencias para formular problemas de datos, diseñar flujos de procesamiento y evaluar soluciones analíticas. Integra fundamentos de sistemas distribuidos con prácticas que pueden realizarse en un portátil, mediante Python, SQL, DuckDB y Apache Spark en modo local.

La organización principal comprende cinco fines de semana, con 56 horas y 15 minutos de formación sincrónica y 7 horas y 45 minutos de trabajo autónomo orientado. Los recesos y el almuerzo están fuera de las 64 horas académicas. Las horas corresponden a períodos de 60 minutos.

# 1 Identificación y alcance
''')
table(plan,['Característica','Definición'],[
('Nombre','Big Data fundamentos ingeniería y analítica aplicada'),('Nivel','Posgrado'),('Modalidad','Online con laboratorios locales y acompañamiento sincrónico'),('Duración','64 horas de trabajo académico'),('Distribución','56 h 15 min sincrónicas y 7 h 45 min autónomas'),('Calendario','Cinco fines de semana consecutivos o según calendario institucional'),('Horario','Viernes 18:00 a 22:00 y sábados 08:00 a 17:00'),('Producto final','Pipeline reproducible, análisis, informe técnico y defensa individual'),('Requisitos','Python básico, SQL elemental y estadística descriptiva')],[4.0,12.8])
text(plan, '''El curso se dirige a estudiantes de posgrado y profesionales que deban analizar datos o evaluar soluciones tecnológicas en investigación y organizaciones. No presupone experiencia en administración de clústeres. Los participantes deben poder ejecutar programas sencillos, interpretar una tabla y construir consultas con filtros y agrupaciones.

La profundidad práctica se concentra en procesamiento por lotes, calidad, SQL analítico, Spark local y un flujo de eventos. Hadoop, NoSQL, lakehouse y despliegue en producción se estudian como decisiones de arquitectura. El aprendizaje automático se utiliza para integrar ingeniería y analítica, con un modelo introductorio; no constituye un curso completo de aprendizaje automático.

# 2 Justificación

El aumento de datos no convierte automáticamente cualquier problema en un problema de Big Data. La selección de una solución requiere conocer los límites de memoria, almacenamiento, latencia, concurrencia y confiabilidad, además del propósito del análisis. Una herramienta distribuida puede añadir complejidad sin mejorar el resultado cuando la carga cabe en un equipo local.

El programa permite contrastar esas decisiones mediante experimentos reproducibles. El estudiante aprende a conservar el significado del dato durante su transformación, medir la calidad, explicar el costo de mover información y distinguir entre evidencia experimental y expectativas de escalabilidad. La exigencia de posgrado se expresa en la argumentación, la evaluación crítica y la documentación de límites.

# 3 Objetivo general

Diseñar, implementar y evaluar una solución reproducible de procesamiento y análisis de datos, utilizando herramientas ejecutables en un portátil y justificando sus decisiones de arquitectura, calidad, rendimiento y gobernanza.

# 4 Resultados de aprendizaje
''')
table(plan,['Código','Resultado verificable','Evidencia'],[
('RA1','Formular un problema y justificar la necesidad de tecnologías de Big Data.','Ficha del problema y comparación de alternativas.'),
('RA2','Diseñar e implementar un flujo de ingesta, calidad y almacenamiento.','Pipeline y reporte de conciliación de registros.'),
('RA3','Construir transformaciones y consultas con Spark y explicar su ejecución.','Laboratorio por lotes y lectura del plan de ejecución.'),
('RA4','Evaluar formatos, configuraciones y procesamiento de eventos con mediciones.','Informe experimental y laboratorio de streaming.'),
('RA5','Integrar analítica, reproducibilidad y gobernanza en una solución defendible.','Proyecto, informe y sustentación individual.')],[2.0,7.8,7.0])
chapter(plan,'5 Organización del tiempo')
table(plan,['Jornada','Tiempo de conexión','Pausas','Tiempo formativo'],[
('Viernes','4 h','15 min','3 h 45 min'),('Sábado','9 h','30 min de recesos y 1 h de almuerzo','7 h 30 min'),('Un fin de semana','13 h','1 h 45 min','11 h 15 min'),('Cinco fines de semana','65 h','8 h 45 min','56 h 15 min'),('Trabajo autónomo','7 h 45 min','No aplica','7 h 45 min'),('Total académico','No incluye pausas','No incluye almuerzos','64 h')],[3.2,3.3,5.6,4.7])
table(plan,['Día','Horario','Actividad'],[
('Viernes','18:00–19:00','Fundamentos y discusión de caso'),('Viernes','19:00–20:00','Demostración y ejercicio guiado'),('Viernes','20:00–20:15','Coffee break'),('Viernes','20:15–21:30','Taller aplicado'),('Viernes','21:30–22:00','Socialización y evaluación formativa'),('Sábado','08:00–09:30','Desarrollo conceptual y demostración'),('Sábado','09:30–09:45','Receso de la mañana'),('Sábado','09:45–12:00','Laboratorio guiado'),('Sábado','12:00–13:00','Almuerzo'),('Sábado','13:00–15:00','Reto práctico y proyecto'),('Sábado','15:00–15:15','Receso de la tarde'),('Sábado','15:15–16:30','Integración y discusión'),('Sábado','16:30–17:00','Retroalimentación y cierre')],[2.2,3.7,10.9])
text(plan,'''En el último sábado los bloques de la tarde se destinan a revisión de resultados y sustentaciones. Para grupos grandes se organizarán equipos de dos o tres personas y una defensa individual breve. Con más de ocho equipos se requerirá distribuir algunas defensas durante la mañana o habilitar salas con evaluadores adicionales, manteniendo la carga horaria.

## Variante de 64 horas sincrónicas

Si la institución exige que las 64 horas sean de interacción sincrónica, se conservarán los cinco fines de semana completos y se añadirán un sexto viernes de 18:00 a 22:00, con 15 minutos de receso, y un sexto sábado de 08:00 a 12:15, con 15 minutos de receso. Se obtienen 3 h 45 min más 4 h, para un total de 64 horas efectivas. El sábado abreviado no incluye almuerzo ni receso de tarde porque termina antes de esos períodos.

En esa variante, las tareas de consolidación y la preparación del informe se realizarán durante las sesiones adicionales, y las defensas finales pasarán al sexto sábado. Si todos los sábados deben finalizar a las 17:00, seis fines de semana sumarían 67 h 30 min efectivos; por tanto, no corresponden exactamente a una carga de 64 horas.
''')
chapter(plan,'6 Estructura curricular')
table(plan,['Módulo','Sincrónico','Autónomo','Total'],[
('1 Fundamentos y problema','11 h 15 min','45 min','12 h'),('2 Arquitecturas almacenamiento y calidad','11 h 15 min','1 h 45 min','13 h'),('3 Procesamiento con Spark','11 h 15 min','1 h 45 min','13 h'),('4 Rendimiento y eventos','11 h 15 min','1 h 45 min','13 h'),('5 Analítica gobernanza e integración','11 h 15 min','1 h 45 min','13 h'),('Total','56 h 15 min','7 h 45 min','64 h')],[7.5,3.1,3.1,3.1])
modules=[
('Módulo 1 Fundamentos y formulación del problema','RA1','Big Data y restricciones de procesamiento; volumen, velocidad, variedad y veracidad; escalabilidad vertical y horizontal; paralelismo; latencia y rendimiento; ciclo de vida del dato y componentes de una solución distribuida.','Diagnóstico, verificación del entorno, generación de datos sintéticos, perfil inicial y definición de pregunta analítica.','Ficha con decisión a apoyar, variables, unidad de análisis, criterios de éxito y restricciones.','45 minutos para precisar el problema y completar el diccionario inicial.'),
('Módulo 2 Arquitecturas almacenamiento y calidad','RA2','Warehouse, lake y lakehouse; papel de Hadoop y HDFS; modelos relacional y NoSQL; CSV, JSON y Parquet; compresión y particionamiento; ETL y ELT; reglas, contratos y cuarentena.','Tipado, deduplicación exacta, clasificación de errores, conciliación y exportación a Parquet con DuckDB.','Pipeline ejecutable, reporte de calidad y datos depurados.','1 h 45 min para documentar reglas, resolver una consulta y redactar una decisión de arquitectura.'),
('Módulo 3 Procesamiento con Apache Spark','RA3','Driver, ejecutores, particiones, trabajos y etapas; evaluación diferida; acciones y transformaciones; DataFrames y Spark SQL; joins, ventanas y shuffle.','Consultas equivalentes en Spark, enriquecimiento con una dimensión y selección de los tres pedidos de mayor importe por ciudad.','Resultados conciliados con DuckDB, código y explicación de plan de ejecución.','1 h 45 min para consolidar el pipeline, añadir una consulta y justificar la estrategia de join.'),
('Módulo 4 Rendimiento y procesamiento de eventos','RA4','Planes de ejecución; caché, particiones y sesgo; benchmarking; throughput y latencia; microbatches; tiempo de evento y procesamiento; ventanas, watermark y checkpoint.','Medición CSV–Parquet y simulación de archivos que llegan por lotes, con eventos tardíos.','Informe de medición y explicación de actualizaciones y descarte de eventos fuera del umbral.','1 h 45 min para interpretar mediciones y documentar un incidente de streaming.'),
('Módulo 5 Analítica gobernanza y proyecto','RA5','Formulación del objetivo analítico; variables y partición temporal; línea base, precisión, recall y F1; fuga de información; visualización; linaje, accesos, privacidad y costo.','Modelo introductorio de entregas tardías, ficha de gobernanza, integración y defensa.','Solución reproducible, informe técnico y sustentación individual.','1 h 45 min antes del último sábado para preparar el informe y la presentación.')]
for title,ra,topics,practice,ev,aut in modules:
    h(plan,title,2);p(plan,f'Resultado principal: {ra}. Contenidos: {topics}');p(plan,'Práctica: '+practice);p(plan,'Evidencia: '+ev);p(plan,'Trabajo autónomo: '+aut)
chapter(plan,'7 Secuencia de sesiones y acompañamiento')
table(plan,['Sesión','Desarrollo central','Producto al cierre'],[
('1 Viernes','Diagnóstico, fundamentos y selección del problema.','Pregunta y criterios de éxito.'),('2 Sábado','Entorno, generación de datos y exploración.','Perfil inicial y diccionario.'),('3 Viernes','Arquitecturas, modelos y formatos.','Decisión inicial de almacenamiento.'),('4 Sábado','ETL, calidad, consultas y Parquet.','Pipeline de limpieza conciliado.'),('5 Viernes','Modelo de ejecución de Spark.','Primeras transformaciones explicadas.'),('6 Sábado','Agregaciones, joins y ventanas.','Resultados equivalentes entre motores.'),('7 Viernes','Medición y optimización.','Diseño del experimento y primeras mediciones.'),('8 Sábado','Streaming, tiempos, ventanas y fallos.','Traza de eventos y análisis de resultados.'),('9 Viernes','Analítica, evaluación y gobernanza.','Modelo base y ficha de gobernanza.'),('10 Sábado','Integración, reproducibilidad y defensa.','Proyecto final y reflexión individual.')],[2.8,7.4,6.6])
text(plan,'''## Estrategia didáctica

Se destinará aproximadamente un 35 % del tiempo sincrónico a fundamentos, discusión y explicación de resultados, y un 65 % a demostraciones, laboratorios y proyecto. Cada encuentro alternará exposición breve, predicción del resultado, ejecución y discusión de evidencia. Los estudiantes trabajarán en parejas para revisar decisiones, pero deberán explicar individualmente su contribución.

El aula virtual contendrá el programa, la guía académica, los archivos de práctica, las consignas y los espacios de entrega. Los notebooks o programas se ejecutarán localmente; la videoconferencia se utilizará para demostraciones y acompañamiento. Las guías escritas permitirán continuar el trabajo cuando haya interrupciones de conexión.

## Apoyo y nivelación

El diagnóstico se aplicará durante la primera sesión y no tendrá nota. Si aparecen dificultades, se usarán ejercicios cortos de lectura de CSV, filtrado, agrupación, joins y estadística descriptiva dentro de los laboratorios. La instalación se verificará durante la primera jornada práctica; no se presupone una actividad obligatoria adicional fuera de las 64 horas.

Cada entrega recibirá retroalimentación sobre corrección, reproducibilidad y razonamiento. El docente facilitará ejemplos mínimos y datos reducidos para resolver fallos de entorno sin convertir toda la sesión en soporte técnico.
''')
chapter(plan,'8 Evaluación y proyecto integrador')
table(plan,['Componente','Peso','Resultados'],[
('Laboratorios de ingesta calidad y procesamiento','30 %','RA1 RA2 RA3'),('Informe experimental de rendimiento y streaming','20 %','RA4'),('Proyecto integrador e informe técnico','35 %','RA1 a RA5'),('Sustentación individual','15 %','RA3 RA4 RA5'),('Total','100 %','')],[10.0,2.3,4.5])
text(plan,'''Los laboratorios se distribuirán en tres entregas de 10 %: problema y perfilado; calidad y almacenamiento; procesamiento con Spark. El informe experimental integrará las prácticas de benchmarking y eventos. El diagnóstico y los ejercicios de cierre tendrán función formativa.

El proyecto resolverá una pregunta sobre pedidos y tiempos de entrega con datos sintéticos, o un problema equivalente autorizado por el docente. El cambio de dominio deberá conservar una complejidad comparable y permitir ejecutar todo el flujo localmente. Se debe evitar el uso de información personal real cuando no sea indispensable.

## Entregables del proyecto

- Ficha del problema con unidad de análisis, pregunta, restricciones y criterios de éxito.
- Diccionario de datos, reglas de calidad y reporte de conciliación.
- Programas o notebooks ejecutables, versiones y orden de ejecución.
- Datos depurados en Parquet y resultados de consultas con Spark.
- Experimento de rendimiento y evidencia de procesamiento de eventos.
- Análisis o modelo sencillo con línea base y evaluación pertinente.
- Informe técnico de 8 a 12 páginas y presentación de 8 a 10 minutos, más preguntas.

## Criterios de valoración

La solución deberá producir resultados correctos, conservar trazabilidad y poder ejecutarse desde un entorno limpio. Las decisiones de arquitectura y rendimiento se evaluarán por su justificación y evidencia, no por utilizar más herramientas. En la analítica se exigirá separar entrenamiento y prueba, declarar limitaciones y evitar conclusiones causales no sustentadas.

La guía académica complementaria contiene rúbricas con descriptores y un banco de preguntas. La escala de aprobación y las reglas institucionales de asistencia se aplicarán según la normativa del programa que ofrezca el curso.
''')
chapter(plan,'9 Recursos y viabilidad técnica')
table(plan,['Recurso','Función'],[('Python y JupyterLab','Programación y documentación de prácticas.'),('pandas y PyArrow','Perfilado y lectura de archivos columnares.'),('DuckDB','SQL analítico local y controles de calidad.'),('PySpark en modo local','Procesamiento por lotes y microbatches.'),('scikit learn y Matplotlib','Modelo introductorio y visualización.'),('Aula virtual y videoconferencia','Distribución de material, acompañamiento y entregas.')],[5.0,11.8])
text(plan,'''Se recomienda un equipo con 16 GB de RAM, cuatro núcleos y 20 a 30 GB libres en SSD. En equipos de 8 GB se utilizarán muestras de menor tamaño y dos hilos de procesamiento. El conjunto inicial tendrá 20 000 registros únicos antes de los errores introducidos y podrá ampliarse gradualmente. No se requiere GPU ni contratar una nube.

La guía propone Python 3.12 y Java 17, con versiones fijadas de las bibliotecas. PySpark 4.0.1 admite instalación local y requiere Java 17 o posterior [R1]. El docente deberá conservar el entorno probado durante la cohorte y verificar la compatibilidad antes de actualizar versiones.

Spark local permite aprender particiones, transformaciones y planes, pero no reproduce redes, fallos de nodos ni escalabilidad de un clúster real. Las conclusiones de rendimiento deberán limitarse al equipo y a las cargas medidos. El caso sintético permite verificar resultados; no valida decisiones comerciales sobre una población real.

## Preparación del docente

- Publicar instrucciones y archivos de práctica en el aula virtual.
- Ejecutar el flujo completo en un equipo comparable al del estudiante.
- Preparar una muestra pequeña y otra ampliada, evitando cargas que bloqueen el portátil.
- Comprobar Java, Python y Spark antes de la sesión correspondiente.
- Distribuir criterios de evaluación y calendario relativo de entregas.
- Reservar tiempos de sustentación según el número de equipos.
''')
h(plan,'10 Bibliografía y documentación')
references(plan)
plan.save(ROOT/'01_Programa_academico_Big_Data_64_horas.docx')

# El manuscrito docente se desarrolla en un módulo independiente para mantener
# separadas la estructura institucional y las explicaciones de aula.
exec((ROOT/'contenido_guia.py').read_text(), globals())
